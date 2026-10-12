import json
import time
import zipfile
from types import SimpleNamespace
from unittest.mock import Mock

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.rag.api import create_app
from app.rag.chunking import build_document_chunks
from app.rag.indexing import build_index, current_generation, safe_extract
from app.rag.parsing import extract_docx_blocks
from app.rag.pipeline import Pipeline, parse_json_response
from app.rag.retrieval import Retriever, build_context


def test_persisted_hybrid_indices_are_independent(corpus):
    settings, embedder = corpus
    retriever = Retriever(settings, embedder)
    assert len(retriever.document_bm25.doc_len) == len(retriever.chunks)
    assert len(retriever.faq_bm25.doc_len) == len(retriever.faq)
    assert retriever.candidates("Библиотека учебники")[0]["document_id"] == "library"
    assert retriever.prepared_answer("Что такое СОП") == "Студенческая оценка преподавания"
    assert retriever.prepared_answer("What is SOP?") is None
    assert retriever.prepared_answer("🙂") is None


def test_index_reused_and_new_generation_on_changed_data(corpus):
    settings, embedder = corpus
    first = current_generation(settings.index_dir)
    assert build_index(settings, embedder) == first
    with zipfile.ZipFile(settings.kb_zip, "a") as archive:
        archive.writestr("note.txt", "updated archive")
    second = build_index(settings, embedder)
    assert second != first
    assert first.is_dir()
    assert current_generation(settings.index_dir) == second


def test_failed_index_keeps_previous_generation(corpus):
    settings, embedder = corpus
    first = current_generation(settings.index_dir)
    settings.kb_zip.write_bytes(b"not a zip file")
    with pytest.raises(zipfile.BadZipFile):
        build_index(settings, embedder)
    assert current_generation(settings.index_dir) == first


def test_optional_faq_and_links(corpus):
    settings, embedder = corpus
    settings.faq_path.unlink()
    build_index(settings, embedder)
    retriever = Retriever(settings, embedder)
    assert retriever.prepared_answer("Что такое СОП") is None
    retriever.links = []
    assert retriever.best_links("Библиотека") == []


def test_faq_threshold_checks_rrf_winner(corpus):
    settings, embedder = corpus
    retriever = Retriever(settings, embedder)
    retriever.settings.faq_rrf_threshold = 0.02
    retriever.faq = [{"answer": "dense winner"}, {"answer": "lexical winner"}]
    retriever.faq_vectors = np.array([[0.99, 0], [0.5, 0]], dtype="float32")
    retriever.encode_query = lambda _: np.array([[1, 0]], dtype="float32")
    retriever.faq_bm25 = SimpleNamespace(get_scores=lambda _: np.array([0, 10]))
    assert retriever.prepared_answer("библиотека") is None


@pytest.mark.parametrize("path", ["../outside.txt", "/outside.txt"])
def test_zip_traversal_rejected(tmp_path, path):
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.writestr(path, "unsafe")
    with pytest.raises(ValueError, match="Unsafe path"):
        safe_extract(archive, tmp_path / "unpacked")


def test_docx_table_and_heading_preserved(tmp_path):
    document = Document()
    document.add_heading("Экзамены", level=1)
    document.add_paragraph("Порядок проведения экзаменов и пересдач. " * 4)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Срок", "Условие"], ["Три дня", "Подать заявление"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / "document.docx"
    document.save(path)
    blocks = extract_docx_blocks(path)
    assert blocks[0]["kind"] == "heading"
    assert blocks[-1]["text"] == "Срок: Три дня | Условие: Подать заявление"
    chunks = build_document_chunks(
        blocks, {"document_id": "exam", "file": path.name, "title": "Экзамены"}
    )
    assert chunks[0]["section"] == "Экзамены"
    assert "Подать заявление" in chunks[0]["text"]


def test_context_budget_and_neighbours(corpus):
    settings, embedder = corpus
    retriever = Retriever(settings, embedder)
    ranked = retriever.candidates("Экзамены")
    selected = retriever.select_context([{**x, "rerank_score": 1.0} for x in ranked])
    assert len(selected) <= 9
    assert len({x["idx"] for x in selected}) == len(selected)
    assert len(build_context(selected, max_chars=500)) <= 500


def test_pipeline_routes_and_language(corpus):
    settings, embedder = corpus
    retriever = Retriever(settings, embedder)

    def response(text, tokens):
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=text))],
            usage=SimpleNamespace(total_tokens=tokens),
        )

    client = Mock()
    client.chat.completions.create.side_effect = [
        response("not json", 20),
        response("Ask the library.", 30),
    ]
    pipeline = Pipeline(settings, retriever, client)
    faq = pipeline.answer("Что такое СОП", [])
    assert faq["route"] == "faq"
    client.chat.completions.create.assert_not_called()
    result = pipeline.answer("Where is the library?", [])
    assert result["route"] == "rag" and result["tokens"] == 50
    final_call = client.chat.completions.create.call_args.kwargs
    assert "английском" in final_call["messages"][0]["content"]
    retriever.settings.link_threshold = 0
    assert retriever.best_links("Библиотека")[0]["url"] == "https://library.hse.ru/"


def test_reranker_ignores_unknown_ids_and_duplicates(corpus):
    settings, embedder = corpus
    retriever = Retriever(settings, embedder)
    candidates = retriever.candidates("Экзамены")
    first = candidates[0]["idx"]
    pipeline = Pipeline(settings, retriever, Mock())
    pipeline.chat = lambda *a, **k: (
        json.dumps(
            {
                "ranked": [
                    {"id": first, "score": 0.9},
                    {"id": first, "score": 0.5},
                    {"id": 999, "score": 1},
                ]
            }
        ),
        7,
    )
    ranked, tokens = pipeline.rerank("Экзамены", candidates, time.perf_counter() + 60)
    assert ranked[0]["rerank_score"] == 0.9
    assert len({x["idx"] for x in ranked}) == len(ranked)
    assert tokens == 7
    assert parse_json_response('```json\n{"ranked": []}\n```') == {"ranked": []}


def test_api_auth_validation_and_upstream_error(corpus):
    settings, _ = corpus
    pipeline = Mock()
    pipeline.answer.return_value = {
        "answer": "Ответ",
        "sources": [],
        "route": "rag",
        "tokens": 1,
        "time_sec": 0.1,
    }
    with TestClient(create_app(settings, pipeline)) as client:
        assert client.get("/health").status_code == 200
        assert client.post("/answer", json={"question": "Экзамен"}).status_code == 401
        pipeline.answer.assert_not_called()
        auth = {"Authorization": "Bearer test-internal-key"}
        for invalid in [
            {"question": " "},
            {"question": "🙂"},
            {"question": "a" * 4001},
            {"question": "Экзамен", "history": ["a"] * 11},
        ]:
            assert client.post("/answer", headers=auth, json=invalid).status_code == 422
        assert (
            client.post("/answer", headers=auth, json={"question": "Экзамен"}).json()["answer"]
            == "Ответ"
        )
        pipeline.answer.side_effect = RuntimeError("private provider details")
        failed = client.post("/answer", headers=auth, json={"question": "Экзамен"})
        assert failed.status_code == 502
        assert "private" not in failed.text
