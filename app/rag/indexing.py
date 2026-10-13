"""Build versioned artifacts once; runtime never unpacks the knowledge base."""

import hashlib
import json
import logging
import os
import re
import stat
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

import faiss
import numpy as np
from openpyxl import load_workbook

from app.config import Settings
from app.rag.chunking import build_document_chunks
from app.rag.parsing import extract_blocks

log = logging.getLogger(__name__)
FORMAT_VERSION = 1


def make_embedder(settings: Settings):
    import torch
    from sentence_transformers import SentenceTransformer

    started = time.perf_counter()
    torch.set_num_threads(settings.embed_threads)
    model = SentenceTransformer(settings.embed_model, device=settings.embed_device)
    if settings.embed_quantize and settings.embed_device == "cpu":
        torch.ao.quantization.quantize_dynamic(
            model, {torch.nn.Linear}, dtype=torch.qint8, inplace=True
        )
    log.info(
        "embedding_model_ready model=%s quantized=%s seconds=%.3f",
        settings.embed_model,
        settings.embed_quantize,
        time.perf_counter() - started,
    )
    return model


def encode_passages(embedder, texts: list[str], batch_size: int) -> np.ndarray:
    return embedder.encode(
        [f"passage: {text}" for text in texts],
        batch_size=batch_size,
        normalize_embeddings=True,
        convert_to_numpy=True,
        show_progress_bar=False,
    ).astype("float32")


def retrieval_text(chunk: dict) -> str:
    return (
        f"Название документа: {chunk['title']}\n"
        f"Категория: {chunk['category']}\n"
        f"Раздел: {chunk['section']}\nТекст: {chunk['text']}"
    )


def safe_extract(archive: Path, destination: Path) -> None:
    root = destination.resolve()
    with zipfile.ZipFile(archive) as zipped:
        total = 0
        for member in zipped.infolist():
            total += member.file_size
            path = (root / member.filename).resolve()
            if not path.is_relative_to(root) or stat.S_ISLNK(member.external_attr >> 16):
                raise ValueError("Unsafe path in knowledge base archive")
            if total > 512 * 1024 * 1024:
                raise ValueError("Knowledge base exceeds 512 MiB unpacked")
        zipped.extractall(root)


def read_chunks(root: Path) -> list[dict]:
    manifests = list(root.rglob("manifest.json"))
    if len(manifests) != 1:
        raise ValueError("Knowledge base must contain exactly one manifest.json")
    manifest_path = manifests[0]
    documents = json.loads(manifest_path.read_text(encoding="utf-8"))["documents"]
    chunks = []
    seen = set()
    for meta in documents:
        if meta["document_id"] in seen:
            raise ValueError("Duplicate document_id in manifest")
        seen.add(meta["document_id"])
        path = (manifest_path.parent / meta["file"]).resolve()
        if not path.is_relative_to(root.resolve()) or not path.is_file():
            raise ValueError(f"Missing or unsafe document: {meta['file']}")
        if path.suffix.lower() not in {".docx", ".pdf", ".txt"}:
            raise ValueError(f"Unsupported document format: {meta['file']}")
        chunks.extend(build_document_chunks(extract_blocks(path), meta))
    if not chunks:
        raise ValueError("No text chunks found in knowledge base")
    return chunks


def read_links(root: Path) -> list[dict]:
    paths = list(root.rglob("all_links.txt"))
    if not paths:
        return []
    if len(paths) > 1:
        raise ValueError("Multiple all_links.txt files in knowledge base")
    links = []
    for line in paths[0].read_text(encoding="utf-8").splitlines():
        match = re.search(r"https?://\S+", line)
        if match:
            links.append(
                {
                    "description": line[: match.start()].strip().rstrip("."),
                    "url": match.group(0),
                }
            )
    return links


def normalize_faq(text: str) -> str:
    return re.sub(r"[^a-zA-Zа-яА-ЯёЁ0-9\s]", "", text).lower().strip()


def read_faq(path: Path) -> list[dict]:
    if not path.is_file():
        log.info("FAQ file absent; answers will use RAG")
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = iter(workbook.active.iter_rows(values_only=True))
        header = next(rows)
        required = ["Основной вопрос", "Варианты вопроса на русском", "Ответ на русском"]
        positions = [header.index(column) for column in required]
        faq = []
        for row_id, row in enumerate(rows):
            question, variants, answer = [row[i] for i in positions]
            if answer is None or not str(answer).strip():
                continue
            for variant in f"{question or ''}\n{variants or ''}".splitlines():
                cleaned = normalize_faq(variant)
                if cleaned:
                    faq.append({"id": row_id, "question": cleaned, "answer": str(answer)})
        return faq
    finally:
        workbook.close()


def fingerprint(settings: Settings) -> str:
    digest = hashlib.sha256()
    digest.update(
        f"{FORMAT_VERSION}:{settings.embed_model}:{settings.embed_quantize and settings.embed_device == 'cpu'}".encode()
    )
    for path in [settings.kb_zip, settings.faq_path]:
        digest.update(path.name.encode())
        if path.is_file():
            with path.open("rb") as stream:
                for block in iter(lambda: stream.read(1024 * 1024), b""):
                    digest.update(block)
        else:
            digest.update(b"missing")
    for name in ["parsing.py", "chunking.py", "indexing.py"]:
        digest.update(Path(__file__).with_name(name).read_bytes())
    return digest.hexdigest()


def current_generation(index_dir: Path) -> Path:
    name = (index_dir / "CURRENT").read_text(encoding="utf-8").strip()
    if not re.fullmatch(r"[a-f0-9]{32}", name):
        raise ValueError("Invalid index generation")
    return index_dir / name


def build_index(settings: Settings, embedder=None) -> Path:
    from app.rag.progress import IndexObserver

    observer = IndexObserver(settings)
    try:
        result = _build_index(settings, embedder, observer)
        observer.report("completed", force=True)
        observer.state.component("indexer", "completed", observer.run_id)
        return result
    except BaseException:
        observer.report("interrupted", force=True)
        observer.state.component("indexer", "interrupted", observer.run_id)
        log.exception("indexing_interrupted")
        raise


def _build_index(settings: Settings, embedder, observer) -> Path:
    if not settings.kb_zip.is_file():
        raise FileNotFoundError(f"Place knowledge base archive at {settings.kb_zip}")
    settings.index_dir.mkdir(parents=True, exist_ok=True)
    build_hash = fingerprint(settings)
    try:
        existing = current_generation(settings.index_dir)
        metadata = json.loads((existing / "metadata.json").read_text(encoding="utf-8"))
        if metadata["build_hash"] == build_hash and all(
            (existing / name).is_file()
            for name in ["documents.faiss", "chunks.json", "links.json", "faq.json", "extras.npz"]
        ):
            observer.files = metadata.get("files", [])
            observer.completed = set(observer.files)
            observer.total_chunks = observer.encoded = metadata["chunks"]
            log.info("Knowledge base unchanged; reusing index %s", existing.name)
            return existing
    except (OSError, ValueError, KeyError):
        pass

    embedder = embedder or make_embedder(settings)
    with tempfile.TemporaryDirectory() as unpacked:
        root = Path(unpacked)
        safe_extract(settings.kb_zip, root)
        chunks = read_chunks(root)
        observer.prepare(root, chunks)
        links = read_links(root)
        faq = read_faq(settings.faq_path)
        vectors = observer.encode(embedder, chunks, settings.embed_batch_size)
        dimension = vectors.shape[1]
        index = faiss.IndexFlatIP(dimension)
        index.add(vectors)
        link_vectors = (
            encode_passages(embedder, [x["description"] for x in links], settings.embed_batch_size)
            if links
            else np.empty((0, dimension), dtype="float32")
        )
        faq_vectors = (
            encode_passages(embedder, [x["question"] for x in faq], settings.embed_batch_size)
            if faq
            else np.empty((0, dimension), dtype="float32")
        )

    generation = uuid.uuid4().hex
    with tempfile.TemporaryDirectory(dir=settings.index_dir) as staging:
        target = Path(staging)
        faiss.write_index(index, str(target / "documents.faiss"))
        for name, payload in [("chunks", chunks), ("links", links), ("faq", faq)]:
            (target / f"{name}.json").write_text(
                json.dumps(payload, ensure_ascii=False), encoding="utf-8"
            )
        np.savez(target / "extras.npz", links=link_vectors, faq=faq_vectors)
        metadata = {
            "format_version": FORMAT_VERSION,
            "build_hash": build_hash,
            "embed_model": settings.embed_model,
            "chunks": len(chunks),
            "links": len(links),
            "faq": len(faq),
            "indexed_files": len({chunk["file"] for chunk in chunks}),
            "total_files": len(observer.files),
            "files": list(dict.fromkeys(chunk["file"] for chunk in chunks)),
            "complete": True,
            "quantized": settings.embed_quantize and settings.embed_device == "cpu",
        }
        (target / "metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
        os.replace(target, settings.index_dir / generation)
    pointer = settings.index_dir / "CURRENT.tmp"
    pointer.write_text(generation, encoding="utf-8")
    os.replace(pointer, settings.index_dir / "CURRENT")
    log.info("Index built: chunks=%d links=%d faq=%d", len(chunks), len(links), len(faq))
    return settings.index_dir / generation


if __name__ == "__main__":
    from app.logging_setup import setup_logging

    settings = Settings()
    setup_logging("indexer", settings)
    build_index(settings)
